from __future__ import annotations

import csv
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "三人分工提交"
IGNORE_DIRS = {"node_modules", "unpackage", ".hbuilderx", "target", ".idea", "dist", "__pycache__"}


PARTS = {
    "成员1_后端与数据库": {
        "role": "后端与数据库",
        "branch": "member1-backend",
        "files": [
            "README.md",
            "backend/pom.xml",
            "backend/src/main/resources/application.yml",
            "tools/smoke_test.ps1",
            "tools/sql_summary.py",
            "tools/inspect_xlsx.py",
            "tools/import_2020_data.ps1",
            "tools/demo-check.ps1",
        ],
        "dirs": [
            "backend/src/main/java",
            "backend/src/main/resources/mapper",
            "backend/src/main/resources/demo-data",
            "sql",
        ],
    },
    "成员2_Web管理端": {
        "role": "Web管理端",
        "branch": "member2-web",
        "files": [
            "README.md",
            "web/index.html",
            "web/package.json",
            "web/package-lock.json",
            "web/vite.config.js",
            "web/README.md",
            "tools/check_web.js",
            "tools/web-dashboard.png",
            "tools/web-report.png",
            "tools/image_palette.py",
        ],
        "dirs": [
            "backend/src/main/webapp",
            "backend/src/main/resources/static",
            "web/src",
        ],
    },
    "成员3_移动端与测试文档": {
        "role": "移动端与测试文档",
        "branch": "member3-mobile-docs",
        "files": [
            "README.md",
            "tools/check_vue.js",
            "tools/mobile_audit.js",
            "tools/build_mobile_audit_docx.py",
            "tools/split_three_submissions.py",
            "tools/setup-adb-reverse.ps1",
            "tools/set-mobile-api.ps1",
            "tools/open-firewall-8006.ps1",
            "tools/ocr_images.ps1",
            "tools/make_github_commits.ps1",
            "tools/current.xml",
        ],
        "dirs": [
            "mobile",
            "docs",
            "demo-files",
            "deliverables/移动端开发阶段检查_Word版",
        ],
    },
}


def safe_clean(path: Path):
    deliverables = (ROOT / "deliverables").resolve()
    resolved = path.resolve()
    if resolved == deliverables or deliverables not in resolved.parents:
        raise RuntimeError(f"Refusing to clean outside deliverables: {resolved}")
    if resolved.exists():
        shutil.rmtree(resolved)
    resolved.mkdir(parents=True, exist_ok=True)


def ignore_callback(directory: str, names: list[str]):
    return {name for name in names if name in IGNORE_DIRS}


def copy_relative(relative: str, part_dir: Path):
    source = ROOT / relative
    if not source.exists():
        raise FileNotFoundError(source)
    destination = part_dir / relative
    destination.parent.mkdir(parents=True, exist_ok=True)
    if source.is_dir():
        shutil.copytree(source, destination, ignore=ignore_callback, dirs_exist_ok=True)
    else:
        shutil.copy2(source, destination)


def write_text(path: Path, content: str):
    path.write_text(content, encoding="utf-8-sig")


def write_submit_readme(part_dir: Path, member: str, role: str, branch: str,
                        files: list[str], dirs: list[str]):
    content = f"""金融财务分析系统 三人分工提交包

提交成员：{member}
负责模块：{role}

一、本包内容
文件：
{chr(10).join('- ' + item for item in files)}

目录：
{chr(10).join('- ' + item for item in dirs)}

二、合并方式
三个成员包的目录结构保持一致，合并时直接把三个包中的 backend、web、mobile、sql、tools、docs、demo-files 目录合并到同一个项目根目录即可。

三、验证命令
成员1：mvn -B -DskipTests package，然后运行 tools/smoke_test.ps1
成员2：浏览器打开 /web/ 或 /console/，使用 root / 123456 登录检查管理端
成员3：node tools/check_vue.js，再用 HBuilderX 运行 mobile 工程

四、说明
本包不包含 node_modules、unpackage、target、.idea 等生成或本机文件。
Git 提交统一使用成员1账号，仓库作者字段不代表个人工作量。

五、Git 分支提交建议
git checkout -b {branch}
git add .
git commit -m "feat: submit {role} part"
git push origin {branch}
三个成员分别推送分支后，由成员1合并到 main。
"""
    write_text(part_dir / "提交说明.txt", content)
    (part_dir / ".gitignore").write_text(
        "\n".join([
            "node_modules/",
            "unpackage/",
            "target/",
            ".idea/",
            "*.log",
            "__pycache__/",
            "",
        ]),
        encoding="utf-8",
    )


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "1F4E79")
        cell._tc.get_or_add_tcPr().append(shade)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        cell.width = Inches(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EAF2F8")
                cell._tc.get_or_add_tcPr().append(shade)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            r.font.name = "Microsoft YaHei"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            r.font.size = Pt(9)
            cell.width = Inches(widths[index])
    return table


def build_division_docx(manifest_rows):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("三人代码分工与提交说明")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(22)
    run.font.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("金融财务分析系统 移动端开发阶段")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_paragraph("本材料将全部需要提交的源码按模块拆成三个互不重叠的提交包。三个包合并后构成完整项目。")
    add_table(
        doc,
        ["成员", "负责模块", "提交包", "建议分支", "贡献占比"],
        [
            ["成员1（LuxCariol）", "后端接口、数据库、数据导入与整体集成", "成员1_后端与数据库.zip", "member1-backend", "40%"],
            ["成员2", "JSP 管理端、Vue Web、静态资源", "成员2_Web管理端.zip", "member2-web", "35%"],
            ["成员3", "uni-app 移动端、测试、文档与验收材料", "成员3_移动端与测试文档.zip", "member3-mobile-docs", "25%"],
            ["合计", "", "", "", "100%"],
        ],
        [1.1, 2.5, 1.45, 1.05, 0.55],
    )
    doc.add_paragraph()
    doc.add_heading("合并方式", level=1)
    doc.add_paragraph("三个提交包使用相同的项目目录结构。将三个包解压到同一个项目根目录，选择合并同名文件夹即可。不要只提交某一个包，否则项目无法独立运行。")
    doc.add_heading("文件范围", level=1)
    doc.add_paragraph(f"文件归属清单共记录 {len(manifest_rows)} 条源文件记录。每个业务文件只归入一个成员包；README.md 作为共享说明同时出现在三个包中。")
    doc.add_heading("提交前检查", level=1)
    for item in [
        "成员1确认后端可编译、SQL 可导入、接口冒烟测试通过。",
        "成员2确认 JSP 管理端和 Vue Web 可登录并完成主要页面操作。",
        "成员3确认移动端模板检查通过、HBuilderX 可运行、真机可访问接口。",
        "三个成员分别提交后，由成员1合并代码并进行最终联调。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
    doc.core_properties.title = "三人代码分工与提交说明"
    doc.core_properties.author = "LuxCariol"
    doc.save(OUT / "三人代码分工说明.docx")


def main():
    safe_clean(OUT)
    manifest_rows = []
    for member, config in PARTS.items():
        part_dir = OUT / member
        part_dir.mkdir(parents=True, exist_ok=True)
        for relative in config["files"]:
            copy_relative(relative, part_dir)
            manifest_rows.append([relative, member, config["role"], "共享说明" if relative == "README.md" else "成员文件"])
        for relative in config["dirs"]:
            copy_relative(relative, part_dir)
            source = ROOT / relative
            for file in source.rglob("*"):
                if not file.is_file():
                    continue
                if any(part in IGNORE_DIRS for part in file.relative_to(ROOT).parts):
                    continue
                manifest_rows.append([
                    file.relative_to(ROOT).as_posix(),
                    member,
                    config["role"],
                    "成员目录",
                ])
        write_submit_readme(part_dir, member, config["role"], config["branch"],
                            config["files"], config["dirs"])

    manifest_rows.sort(key=lambda row: (row[1], row[0]))
    with (OUT / "文件归属清单.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["文件路径", "归属成员", "负责模块", "说明"])
        writer.writerows(manifest_rows)

    build_division_docx(manifest_rows)

    for member in PARTS:
        part_dir = OUT / member
        archive = OUT / member
        shutil.make_archive(str(archive), "zip", root_dir=part_dir)

    overall = OUT / "三人分工提交材料"
    if overall.exists():
        shutil.rmtree(overall)
    overall.mkdir()
    for file in OUT.iterdir():
        if file.name == "三人分工提交材料":
            continue
        if file.is_file():
            shutil.copy2(file, overall / file.name)

    shutil.make_archive(str(OUT / "三人分工提交材料"), "zip", root_dir=overall)
    print(f"Generated three submission packages in {OUT}")


if __name__ == "__main__":
    main()
